# Copyright (c) 2026 Davide Di Filippo. Tutti i diritti riservati. Vedi LICENSE.
"""Crea i file di prova in tutti i formati (serve LibreOffice e pandoc). Uso: python3 tests/make-fixtures.py"""
import os, subprocess, shutil
from pathlib import Path
from docx import Document
from docx.shared import Pt
import openpyxl

HERE = Path(__file__).parent / 'fixtures'
src = (HERE / 'guerra-fredda.txt').read_text(encoding='utf-8').strip().split('\n\n')
title, paras = src[0], src[1:]

# DOCX ricco: titolo, sottotitoli, grassetti, elenco, tabella
d = Document()
d.core_properties.title = title
d.add_heading(title, level=0)
for i, p in enumerate(paras):
    if i == 2:
        d.add_heading('La Germania divisa', level=1)
    if i == 5:
        d.add_heading('Armi nucleari e crisi', level=2)
    par = d.add_paragraph()
    first, _, rest = p.partition(' ')
    r = par.add_run(first + ' '); r.bold = True
    par.add_run(rest)
d.add_heading('Date da ricordare', level=1)
for item in ['1945: fine della seconda guerra mondiale', '1949: nasce la NATO', '1989: cade il muro di Berlino']:
    d.add_paragraph(item, style='List Bullet')
t = d.add_table(rows=3, cols=2)
for r, (a, b) in enumerate([('Blocco', 'Alleanza'), ('Occidentale', 'NATO'), ('Orientale', 'Patto di Varsavia')]):
    t.cell(r, 0).text = a; t.cell(r, 1).text = b
d.save(HERE / 'sample.docx')

# XLSX
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Cronologia'
ws.append(['Anno', 'Evento'])
for a, e in [(1946, 'Discorso di Churchill sulla cortina di ferro'), (1948, 'Blocco di Berlino e ponte aereo'), (1949, 'Nasce la NATO'), (1955, 'Patto di Varsavia'), (1961, 'Costruzione del muro di Berlino'), (1962, 'Crisi dei missili di Cuba'), (1989, 'Caduta del muro di Berlino'), (1991, "Fine dell'Unione Sovietica")]:
    ws.append([a, e])
wb.save(HERE / 'sample.xlsx')

# PPTX
try:
    from pptx import Presentation
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = title; s.placeholders[1].text = 'Riassunto per la lezione'
    for i in range(0, len(paras), 2):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = 'Parte ' + str(i // 2 + 1)
        s.placeholders[1].text = paras[i]
    pr.save(HERE / 'sample.pptx')
except ImportError:
    pass

def soffice(src, fmt, outname=None):
    subprocess.run(['soffice', '--headless', '--convert-to', fmt, '--outdir', str(HERE), str(src)], check=True, capture_output=True, timeout=180)

soffice(HERE / 'sample.docx', 'doc')
soffice(HERE / 'sample.docx', 'odt')
soffice(HERE / 'sample.docx', 'rtf')
soffice(HERE / 'sample.docx', 'pdf')
soffice(HERE / 'sample.docx', 'html')
soffice(HERE / 'sample.xlsx', 'ods')
if (HERE / 'sample.pptx').exists():
    soffice(HERE / 'sample.pptx', 'odp')
subprocess.run(['pandoc', str(HERE / 'sample.docx'), '-o', str(HERE / 'sample.epub'), '--metadata', 'title=' + title], check=True)
subprocess.run(['pandoc', str(HERE / 'sample.docx'), '-t', 'gfm', '-o', str(HERE / 'sample.md')], check=True)
(HERE / 'sample.srt').write_text('1\n00:00:01,000 --> 00:00:04,000\nBenvenuti alla lezione sulla guerra fredda.\n\n2\n00:00:04,500 --> 00:00:08,000\nOggi parliamo di <i>Berlino</i> e della NATO.\n\n3\n00:00:15,000 --> 00:00:19,000\nNel 1989 cadde il muro di Berlino.\n', encoding='utf-8')
(HERE / 'sample.csv').write_text('Anno;Evento\n1949;Nasce la NATO\n1961;Costruzione del muro di Berlino\n1989;Caduta del muro di Berlino\n', encoding='utf-8')
(HERE / 'sample-cp1252.txt').write_bytes('La città di Berlino fu divisa. Perché? Perché era il simbolo della guerra fredda.\n'.encode('cp1252'))
print('\n'.join(sorted(p.name for p in HERE.iterdir())))
